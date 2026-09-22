#!/usr/bin/env python3
"""
import_xlsx.py: turn the catalog workbook into the CSV the virtual library reads.

    python3 import_xlsx.py /path/to/catalog.xlsx            # writes data/my-library.csv
    python3 import_xlsx.py /path/to/catalog.xlsx out.csv    # writes somewhere else

What it does:
  1. reads the "Library" sheet (one row per item, all its columns are kept
     with their original headers),
  2. joins the "Descriptions" sheet on ID and adds three columns: Year (the
     free text, for example "1989 (4th ed.)"), YearParsed (the first four
     digit number in that text, or blank) and Description,
  3. joins the "Shelves & photos" sheet on Unit and Shelf and adds a
     "Shelf description" column, which the shelf view shows under each shelf,
  4. adds an empty "sort_title" column. Fill a cell to change where a book
     files in the alphabetical check without touching its title (for example
     "Arceus and the Jewel of Life" for "Pokémon: Arceus and the Jewel of
     Life"). The virtual library reads it when it is not empty.

The output is UTF-8 CSV with a header row. data/my-library.csv is listed in
.gitignore because the workbook carries personal notes; the sample
data/library.csv in the same schema is what gets committed.

Needs openpyxl (pip install openpyxl). Only the standard library otherwise.
"""
import csv
import os
import re
import sys

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is needed: pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_OUT = os.path.join(HERE, "data", "my-library.csv")

LIBRARY_SHEET = "Library"
DESCRIPTIONS_SHEET = "Descriptions"
SHELVES_SHEET = "Shelves & photos"
EXTRA_COLUMNS = ["Year", "YearParsed", "Description", "Shelf description", "sort_title"]

YEAR_RE = re.compile(r"(?<!\d)(\d{4})(?!\d)")


def text(value):
    """Cell value as a trimmed string; None becomes ''."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_year(value):
    """First four digit number in the free text year, or ''."""
    m = YEAR_RE.search(text(value))
    return m.group(1) if m else ""


def read_sheet(wb, name):
    """Rows of a sheet as dicts keyed by the header row. Blank rows are skipped."""
    if name not in wb.sheetnames:
        return []
    rows = wb[name].iter_rows(values_only=True)
    header = [text(h) for h in next(rows, [])]
    out = []
    for row in rows:
        if row is None or all(v is None or text(v) == "" for v in row):
            continue
        out.append({header[i]: text(v) for i, v in enumerate(row) if i < len(header) and header[i]})
    return out


def main(argv):
    if len(argv) < 2:
        sys.exit(__doc__)
    src = argv[1]
    out_path = argv[2] if len(argv) > 2 else DEFAULT_OUT

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    library = read_sheet(wb, LIBRARY_SHEET)
    if not library:
        sys.exit(f'No rows found on the "{LIBRARY_SHEET}" sheet of {src}.')
    descriptions = {r.get("ID", ""): r for r in read_sheet(wb, DESCRIPTIONS_SHEET) if r.get("ID")}
    shelves = {(r.get("Unit", ""), r.get("Shelf", "")): r.get("Description", "") for r in read_sheet(wb, SHELVES_SHEET)}

    library_columns = list(library[0].keys())
    columns = library_columns + [c for c in EXTRA_COLUMNS if c not in library_columns]

    matched = 0
    with_year = 0
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in library:
            desc = descriptions.get(row.get("ID", ""))
            year_text = ""
            if desc:
                matched += 1
                # The Descriptions sheet names the column with a note in it.
                year_key = next((k for k in desc if k.lower().startswith("year")), None)
                year_text = desc.get(year_key, "") if year_key else ""
                row["Description"] = desc.get("Description", "")
            row["Year"] = year_text
            row["YearParsed"] = parse_year(year_text)
            if row["YearParsed"]:
                with_year += 1
            row["Shelf description"] = shelves.get((row.get("Unit", ""), row.get("Shelf", "")), "")
            row.setdefault("sort_title", "")
            writer.writerow(row)

    shelf_count = len({(r.get("Unit", ""), r.get("Shelf", "")) for r in library})
    print(f"Wrote {len(library)} rows, {shelf_count} shelves, to {out_path}")
    print(f"Descriptions matched for {matched} rows, a parsed year for {with_year}.")


if __name__ == "__main__":
    main(sys.argv)
